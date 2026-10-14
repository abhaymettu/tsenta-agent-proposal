"""Build the Tsenta pitch deck as an editable PPTX (native text, shapes, connectors, speaker notes).

Run: uv run --no-project --with python-pptx python build/build_deck.py
Content lives in build/content.py; this file is layout only.
"""
import copy
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
import content as C  # noqa: E402

INK = RGBColor(0x25, 0x1F, 0x21)
SEC = RGBColor(0x58, 0x52, 0x54)
TEAL = RGBColor(0x73, 0xA8, 0x9A)
TEAL_DK = RGBColor(0x3F, 0x76, 0x68)  # darker teal for small text on white (contrast)
TEAL_TINT = RGBColor(0xE6, 0xF0, 0xED)
RULE = RGBColor(0xE4, 0xE0, 0xDD)
PAPER = RGBColor(0xF6, 0xF4, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

W, H = Inches(13.333), Inches(7.5)
LM = Inches(0.8)  # left margin
CW = W - 2 * LM  # content width


# ---------- low-level helpers ----------

def _run_props(run, size, bold=False, color=INK, italic=False, spacing=None):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    if spacing is not None:
        run._r.get_or_add_rPr().set("spc", str(spacing))


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, name=None):
    """paras: list of dicts {runs:[(text, {opts})] | text, size, bold, color, align, space_after, line}"""
    box = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        if "space_after" in p:
            para.space_after = Pt(p["space_after"])
        if "space_before" in p:
            para.space_before = Pt(p["space_before"])
        if "line" in p:
            para.line_spacing = p["line"]
        runs = p.get("runs") or [(p["text"], {})]
        for t, o in runs:
            r = para.add_run()
            r.text = t
            _run_props(
                r,
                o.get("size", p.get("size", 18)),
                o.get("bold", p.get("bold", False)),
                o.get("color", p.get("color", INK)),
                o.get("italic", p.get("italic", False)),
                o.get("spacing", p.get("spacing")),
            )
    return box


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, dash=False, rounded=False, name=None, radius=0.08):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, int(x), int(y), int(w), int(h))
    if name:
        shp.name = name
    if rounded:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
        if dash:
            ln = shp.line._get_or_add_ln()
            d = etree.SubElement(ln, qn("a:prstDash"))
            d.set("val", "dash")
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = Inches(0.18)
    tf.margin_top = tf.margin_bottom = Inches(0.12)
    return shp


def oval(slide, x, y, w, h, fill=None, line=None, line_w=1.0, dash=False, name=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, int(x), int(y), int(w), int(h))
    if name:
        shp.name = name
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
        if dash:
            ln = shp.line._get_or_add_ln()
            etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    shp.shadow.inherit = False
    return shp


def shape_text(shp, paras, anchor=MSO_ANCHOR.TOP):
    tf = shp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        if "space_after" in p:
            para.space_after = Pt(p["space_after"])
        if "line" in p:
            para.line_spacing = p["line"]
        runs = p.get("runs") or [(p["text"], {})]
        for t, o in runs:
            r = para.add_run()
            r.text = t
            _run_props(r, o.get("size", p.get("size", 16)), o.get("bold", p.get("bold", False)),
                       o.get("color", p.get("color", INK)), o.get("italic", False), o.get("spacing", p.get("spacing")))


def line(slide, x1, y1, x2, y2, color=SEC, width=1.25, arrow=False, dash=False, elbow=False, name=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT, int(x1), int(y1), int(x2), int(y2))
    if name:
        c.name = name
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    if arrow:
        t = etree.SubElement(ln, qn("a:tailEnd"))
        t.set("type", "triangle")
        t.set("w", "med")
        t.set("len", "med")
    return c


def hrule(slide, x, y, w, color=RULE, width=1.0):
    return line(slide, x, y, x + w, y, color=color, width=width)


# ---------- slide furniture ----------

def kicker(slide, label, y=Inches(0.62)):
    text(slide, LM, y, CW, Inches(0.3),
         [{"text": label.upper(), "size": 12, "bold": True, "color": TEAL_DK, "spacing": 120}], name="Kicker")


def title(slide, t, y=Inches(0.98), size=38, h=Inches(1.2), w=None):
    ph = slide.shapes.title
    ph.left, ph.top, ph.width, ph.height = int(LM), int(y), int(w or CW), int(h)
    tf = ph.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.text = ""
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.line_spacing = 0.95
    r = p.add_run()
    r.text = t
    _run_props(r, size, True, INK)
    return ph


STATUS_LINES = {
    "vision": "Vision and roadmap. Only the job agent is live today.",
    "proposed": "Proposal, not built. Only the job agent is live today.",
    "synthetic": "Synthetic concept. Only the job agent is live today.",
    "agenttest": "Real agent on dummy records. Full demo gate not passed.",
}


def footer(slide, n, source=None, status=None):
    hrule(slide, LM, Inches(6.88), CW)
    if status:
        runs = [(STATUS_LINES[status], {"bold": True, "color": INK}), ("   Tsenta  ·  October 14, 2026", {"color": SEC})]
    else:
        runs = [("Tsenta  ·  proposal for discussion  ·  October 14, 2026", {"color": SEC})]
    text(slide, LM, Inches(6.98), Inches(6.6), Inches(0.3), [{"runs": runs, "size": 10}], name="Footer")
    text(slide, W - LM - Inches(0.6), Inches(6.98), Inches(0.6), Inches(0.3),
         [{"text": str(n), "size": 10, "color": SEC, "align": PP_ALIGN.RIGHT}], name="Slide number")
    if source:
        text(slide, Inches(7.5), Inches(6.98), W - LM - Inches(0.7) - Inches(7.5), Inches(0.3),
             [{"text": source, "size": 9, "color": SEC, "align": PP_ALIGN.RIGHT}], name="Source line")


def tag(slide, x, y, label, kind="proposed", w=None):
    """Small status tag: today / proposed / vision / synthetic."""
    styles = {
        "today": (INK, WHITE, None),
        "proposed": (WHITE, TEAL_DK, TEAL),
        "vision": (WHITE, SEC, SEC),
        "synthetic": (PAPER, SEC, None),
    }
    fill, color, ln = styles[kind]
    w = w or Inches(0.22 + 0.09 * len(label))
    shp = rect(slide, x, y, w, Inches(0.27), fill=fill, line=ln, line_w=0.75, rounded=True, radius=0.5, name=f"Tag {label}")
    shp.text_frame.margin_left = shp.text_frame.margin_right = Inches(0.05)
    shp.text_frame.margin_top = shp.text_frame.margin_bottom = 0
    shp.text_frame.word_wrap = False  # a status pill must never break mid-word in any renderer
    shape_text(shp, [{"text": label.upper(), "size": 8.5, "bold": True, "color": color, "align": PP_ALIGN.CENTER, "spacing": 80}],
               anchor=MSO_ANCHOR.MIDDLE)
    return shp


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body.strip()


# ---------- theme and master ----------

def patch_theme(prs):
    theme_part = prs.slide_master.part.part_related_by(RT.THEME)
    xml = theme_part.blob.decode("utf-8")
    scheme = (
        '<a:clrScheme name="Tsenta editorial">'
        '<a:dk1><a:srgbClr val="251F21"/></a:dk1><a:lt1><a:srgbClr val="FFFFFF"/></a:lt1>'
        '<a:dk2><a:srgbClr val="585254"/></a:dk2><a:lt2><a:srgbClr val="F6F4F2"/></a:lt2>'
        '<a:accent1><a:srgbClr val="73A89A"/></a:accent1><a:accent2><a:srgbClr val="3F7668"/></a:accent2>'
        '<a:accent3><a:srgbClr val="585254"/></a:accent3><a:accent4><a:srgbClr val="A9A3A5"/></a:accent4>'
        '<a:accent5><a:srgbClr val="E6F0ED"/></a:accent5><a:accent6><a:srgbClr val="251F21"/></a:accent6>'
        '<a:hlink><a:srgbClr val="3F7668"/></a:hlink><a:folHlink><a:srgbClr val="585254"/></a:folHlink>'
        "</a:clrScheme>"
    )
    xml = re.sub(r"<a:clrScheme .*?</a:clrScheme>", scheme, xml, count=1, flags=re.S)
    xml = re.sub(r'(<a:majorFont>\s*<a:latin typeface=")[^"]*(")', r"\1Arial\2", xml)
    xml = re.sub(r'(<a:minorFont>\s*<a:latin typeface=")[^"]*(")', r"\1Arial\2", xml)
    theme_part._blob = xml.encode("utf-8")


def widen_layouts(prs, factor):
    """Default template is 4:3; stretch master and layout placeholders to the 16:9 canvas."""
    for coll in [prs.slide_master.shapes] + [l.shapes for l in prs.slide_layouts]:
        for shp in coll:
            if shp.left is not None and shp.width is not None:
                shp.left = int(shp.left * factor)
                shp.width = int(shp.width * factor)


# ---------- build ----------

def build(out_path, include_roadmap=True, include_cost=True):
    prs = Presentation()
    old_w = prs.slide_width
    prs.slide_width, prs.slide_height = W, H
    widen_layouts(prs, W / old_w)
    patch_theme(prs)
    prs.core_properties.title = "Tsenta: starts with your job, grows into your life"
    prs.core_properties.author = "Abhay Mettu"
    prs.core_properties.subject = "Vision and proposed architecture for discussion, October 14, 2026"
    layout = prs.slide_layouts[5]  # Title Only

    n = 0

    def new(t, kick=None, tsize=38, ty=Inches(0.98), th=Inches(1.2), tw=None):
        nonlocal n
        n += 1
        s = prs.slides.add_slide(layout)
        if kick:
            kicker(s, kick)
        title(s, t, y=ty, size=tsize, h=th, w=tw)
        return s, n

    for spec in C.slides(include_roadmap=include_roadmap, include_cost=include_cost):
        s, num = new(spec["title"], spec.get("kicker"), spec.get("tsize", 38), spec.get("ty", Inches(0.98)),
                     spec.get("th", Inches(1.2)), spec.get("tw"))
        spec["draw"](s, num, globals())
        if spec.get("footer", True):
            footer(s, num, spec.get("source"), spec.get("status"))
        notes(s, spec["notes"])

    prs.save(out_path)
    return n


if __name__ == "__main__":
    out = ROOT / "outputs" / "Tsenta-Pitch-2026-10-14.pptx"
    roadmap = "--no-roadmap" not in sys.argv
    cost = "--no-cost" not in sys.argv
    count = build(out, include_roadmap=roadmap, include_cost=cost)
    print(f"wrote {out} ({count} slides)")
